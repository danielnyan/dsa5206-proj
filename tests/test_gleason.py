from __future__ import annotations
import ast
import csv
import json
from pathlib import Path
import zipfile

import numpy as np
import pytest
from PIL import Image
from openpyxl import Workbook

from modern_pca.gleason import CLASSES,SCHEDULE,legacy_score,sha,read_rows,write_json,write_csv
from modern_pca.prepare_gleason import crowd_labels,sicap_labels,audit_splits,extract,image_index,tumour_rows
from modern_pca.gleason_regions import summarize,stability
from modern_pca.gleason_wsi import eight_views


def test_gleason2019_pixel_votes_background_ties_and_invalid_six():
    from modern_pca.prepare_gleason import consensus_mask
    masks = [np.array([[3, 3, 0, 6, 1]], dtype=np.uint8),
             np.array([[3, 4, 5, 3, 1]], dtype=np.uint8),
             np.array([[4, 0, 0, 3, 0]], dtype=np.uint8)]
    result, unresolved = consensus_mask(masks)
    assert result.tolist() == [[3, 0, 0, 0, 1]]
    assert unresolved.tolist() == [[False, True, False, True, False]]


def test_gleason2019_background_and_unknown_encoding():
    from modern_pca.prepare_gleason import consensus_mask
    result, unresolved = consensus_mask([np.zeros((2, 2), dtype=np.uint8)])
    assert not result.any() and not unresolved.any()
    with pytest.raises(ValueError, match='encoding'):
        consensus_mask([np.array([[2]], dtype=np.uint8)])


def test_gleason2019_tile_geometry_uses_existing_preprocessing():
    from modern_pca.gleason_regions import preprocess_image
    class Identity:
        def transform(self, value):
            return value
    source = Image.new('RGB', (1200, 600), (240, 10, 70))
    tile = source.crop((600, 0, 1200, 600))
    pixels = preprocess_image(tile, (Identity(), Identity()))
    assert tile.size == (600, 600)
    assert pixels.shape == (350, 350, 3)
    assert np.allclose(pixels[0, 0], np.array([240, 10, 70]) / 255)


@pytest.mark.parametrize('audit_only', [True, False])
def test_gleason2019_preparation_inventory_and_audit(tmp_path, monkeypatch, audit_only):
    import io
    from types import SimpleNamespace
    import modern_pca.prepare_gleason as module
    raw = tmp_path / 'raw'
    images = raw / 'images'
    images.mkdir(parents=True)
    Image.new('RGB', (1250, 600), (240, 10, 70)).save(images / 'slide001_core001.jpg')
    mask = np.ones((600, 1250), dtype=np.uint8)
    mask[:, :600] = 3
    encoded = io.BytesIO()
    Image.fromarray(mask).save(encoded, format='PNG')
    archive = raw / 'masks.zip'
    with zipfile.ZipFile(archive, 'w') as outer:
        for expert in range(1, 7):
            nested = io.BytesIO()
            with zipfile.ZipFile(nested, 'w') as inner:
                inner.writestr('slide001_core001_classimg_nonconvex.png', encoded.getvalue())
            outer.writestr(f'Maps{expert}_T.zip', nested.getvalue())
    monkeypatch.setattr(module, 'GLEASON_MASK_SHA', sha(archive))
    monkeypatch.setattr(module, 'download_gleason2019', lambda root: (images, archive))
    class Identity:
        def transform(self, value):
            return value
    monkeypatch.setattr('modern_pca.evaluate_paper.create_legacy_normalizer',
                        lambda path: (Identity(), Identity()))
    output = tmp_path / 'prepared'
    module.prepare_gleason2019(SimpleNamespace(gleason2019_root=raw, output=output,
        stain_reference=module.REFERENCE, audit_only=audit_only))
    result = json.loads((output / 'gleason2019.json').read_text())
    assert result['complete'] and not result['training_ready']
    assert result['counts'] == {'single_gp': 1, 'non_tumour': 1, 'discarded_edge_pixels': 30000}
    with (output / 'patch_inventory.csv').open() as stream:
        rows = list(csv.DictReader(stream))
    assert [r['x'] for r in rows] == ['0', '600']
    assert rows[0]['class_name'] == 'GP3' and rows[0]['experts'] == '6'
    assert all(r['split'] == 'unassigned' for r in rows)
    assert not rows[1]['cache_file']
    if audit_only:
        assert not rows[0]['cache_file']
    else:
        with Image.open(output / 'cache' / rows[0]['cache_file']) as cached:
            assert cached.size == (350, 350)
        assert sha(output / 'cache' / rows[0]['cache_file']) == rows[0]['cache_sha256']


def test_crowd_uses_supplied_vote_not_new_tie_rule(tmp_path):
    for split in ('train','val','test'):
        key='ground truth' if split=='test' else 'MV'
        row={f'marker{i}':v for i,v in enumerate([0,2,-1,-1,-1,-1,-1],1)}
        row.update({'Patch filename':'case_row_0_column_0',key:2 if split!='test' else 1})
        write_csv(tmp_path/f'{split}.csv',[row])
    rows=list(crowd_labels(tmp_path))
    assert [r['source_label'] for r in rows]==[2,2,1]
    assert [r['split'] for r in rows]==['train','validation','test']
    assert rows[0]['source_wsi']=='case'


def test_sicap_mapping_and_cribriform_is_not_class(tmp_path):
    for name in ('Train','Val','Test'):
        wb=Workbook()
        sheet=wb.active
        sheet.append(['image_name','NC','G3','G4','G5','G4C'])
        sheet.append(['slide_Block_Region_1.jpg',0,0,1,0,1])
        wb.save(tmp_path/f'{name}.xlsx')
    assert [r['source_label'] for r in sicap_labels(tmp_path)]==[2,2,2]
    assert all(r['source_g4c']==1 for r in sicap_labels(tmp_path))
    wb=Workbook(); sheet=wb.active
    sheet.append(['image_name','NC','G3','G4','G5'])
    sheet.append(['bad.jpg',0,1,1,0]); wb.save(tmp_path/'Train.xlsx')
    with pytest.raises(ValueError,match='one-hot'): list(sicap_labels(tmp_path))


@pytest.mark.parametrize('dataset',['crowd','sicap'])
def test_only_nc_removed_and_classes_remapped(dataset):
    records=[dict(dataset=dataset,source_label=k,votes='[-1, -1, 1]') for k in range(4)]
    result=tumour_rows(records)
    assert [r['source_label'] for r in result]==[1,2,3]
    assert [r['label'] for r in result]==[0,1,2]
    assert [r['class_name'] for r in result]==CLASSES
    assert all(r['votes']=='[-1, -1, 1]' for r in result)


def test_masks_not_indexed(tmp_path):
    (tmp_path/'images').mkdir(); (tmp_path/'masks').mkdir()
    for folder in ('images','masks'):
        Image.new('RGB',(4,4)).save(tmp_path/f'{folder}/a.jpg')
    assert image_index(tmp_path)=={'a':tmp_path/'images/a.jpg'}


def test_split_audit_reports_slide_overlap_and_rejects_content_leak():
    rows=[dict(dataset='crowd',name='a',source_wsi='same',split='train',file_sha256='1'),
          dict(dataset='crowd',name='b',source_wsi='same',split='validation',file_sha256='2')]
    assert len(audit_splits(rows)['slide_overlaps'])==1
    rows[1]['file_sha256']='1'
    with pytest.raises(ValueError,match='Identical'): audit_splits(rows)


def test_zip_traversal_rejected(tmp_path):
    path=tmp_path/'other.zip'
    with zipfile.ZipFile(path,'w') as z: z.writestr('../outside','bad')
    with pytest.raises(ValueError,match='Unsafe'): extract(path,tmp_path/'extracted')
    assert not (tmp_path/'outside').exists()


@pytest.mark.parametrize('values,expected',[
    ([96,4,0],(3,3,1)),([70,30,0],(3,4,2)),([30,70,0],(4,3,3)),
    ([0,92,8],(4,5,5)),([95,5,0],(3,3,1)),([94.9,5.1,0],(3,4,2)),
    ([90,5,5],(3,3,1)),([89.9,5,5.1],(3,5,4)),([50,50,0],(4,3,3))])
def test_legacy_scoring(values,expected):
    result=legacy_score(values)
    assert (result['primary'],result['secondary'],result['isup'])==expected


def test_soft_aggregation_does_not_argmax():
    result=summarize([[.6,.4,0],[.7,.3,0]])
    assert result['p_gp4']==pytest.approx(35)
    assert result['isup']==2


def test_schedule_and_propagated_batchnorm():
    assert sum(s[1] for s in SCHEDULE)==14
    assert [s[0].split('_')[-1] for s in SCHEDULE]==['17','14','11','9','7','5','3','1']
    source=(Path(__file__).parents[1]/'modern_pca/train_gleason.py').read_text()
    ast.parse(source)
    assert 'backbone_training=None' in source


def test_eight_unique_d4_views():
    image=Image.fromarray(np.arange(27,dtype=np.uint8).reshape(3,3,3))
    assert len({np.asarray(v).tobytes() for v in eight_views(image)})==8


def test_stability_bounds_sample_count(tmp_path):
    from argparse import Namespace
    records=[dict(image='region',p_gp3=.7,p_gp4=.3,p_gp5=0) for _ in range(16)]
    path=tmp_path/'predictions.csv'; write_csv(path,records)
    stability(Namespace(predictions=path,output=tmp_path/'out',seed=42,
                        rounds=2,max_tiles=19,fov_um=150))
    with (tmp_path/'out/stability.csv').open() as stream: rows=list(csv.DictReader(stream))
    assert len(rows)==32 and max(int(r['tiles']) for r in rows)==16
    assert all(r['same_isup']=='1' for r in rows)


def test_manifest_integrity(tmp_path):
    root=tmp_path/'manifests'; root.mkdir()
    path=root/'crowd_train.csv'
    row=dict(dataset='crowd',split='train',label=0,class_name=CLASSES[0])
    write_csv(path,[row]); write_json(tmp_path/'prepared.json',dict(manifests={path.name:sha(path)}))
    assert read_rows(root,'crowd','train')[0]['label']==0
    path.write_text('tampered')
    with pytest.raises(ValueError,match='checksum'): read_rows(root,'crowd','train')


def test_real_normalization_cache_matches_existing_adapter(tmp_path):
    pytest.importorskip('staintools')
    from modern_pca.prepare_gleason import init_worker,prepare_image
    from modern_pca.gleason import REFERENCE,ROOT
    from modern_pca.evaluate_paper import create_legacy_normalizer,preprocess_legacy
    source=next((ROOT/'1_training/Examples/training_gleason_grading/gp3').glob('*.jpg'))
    row=dict(sample_id='test:a',file_sha256=sha(source),path=str(source))
    init_worker(str(REFERENCE),str(tmp_path))
    cached=prepare_image(row)
    values=np.asarray(Image.open(tmp_path/cached['cache_file']),dtype=np.float32)/255
    expected=preprocess_legacy(row,source,create_legacy_normalizer(REFERENCE))
    np.testing.assert_array_equal(values,expected)


def test_sicap_slide_metadata_and_cohort_roles(tmp_path):
    from modern_pca.prepare_gleason import slide_grades,assign_region_role
    from modern_pca.gleason import select_training
    wb=Workbook(); sheet=wb.active
    sheet.append(['slide_id','patient_id','Gleason_primary','Gleason_secondary'])
    sheet.append(['pure','p1',4,4]); sheet.append(['mixed','p2',3,4])
    path=tmp_path/'wsi_labels.xlsx'; wb.save(path)
    grades=slide_grades(path)
    def row(slide,label=2,dataset='sicap'):
        return dict(dataset=dataset,source_wsi=slide,source_label=label,
                    label=label-1,split='train')
    pure=assign_region_role(row('pure'),grades)
    mixed=assign_region_role(row('mixed'),grades)
    conflict=assign_region_role(row('pure',1),grades)
    unknown=assign_region_role(row('absent'),grades)
    crowd=assign_region_role(row('pure',dataset='crowd'),grades)
    assert [r['region_role'] for r in (pure,mixed,conflict,unknown,crowd)]==[
        'pure','mixed','conflict','unknown','unknown']
    assert pure['patient_id']=='p1'
    rows=[pure,mixed,conflict,unknown,crowd]
    assert select_training(rows,'pure')==[pure]
    assert select_training(rows,'published')==rows
    with pytest.raises(ValueError,match='Only training'):
        select_training([dict(pure,split='test')],'pure')


@pytest.mark.parametrize('label',[0,1,2])
def test_mining_strict_threshold_all_classes(label):
    from modern_pca.gleason import mining_decision
    p=np.full(3,.025); p[label]=.95
    assert mining_decision(p) is None
    p=np.full(3,.0249995); p[label]=.950001
    assert mining_decision(p)==label


@pytest.mark.parametrize('p',[[.5,.4,.1],[.949,.03,.021],[1,1,0],[float('nan'),0,0],
                              [float('inf'),0,0],[-.1,.1,1],[.97,.01],[1.01,0,0]])
def test_mining_invalid_or_unconfident_probabilities(p):
    from modern_pca.gleason import mining_decision
    if len(p)==3 and np.isfinite(p).all() and min(p)>=0 and sum(p)==1:
        assert mining_decision(p) is None
    else:
        with pytest.raises(ValueError): mining_decision(p)


@pytest.mark.parametrize('key',['sample_id','file_sha256','source_wsi','patient_id'])
def test_mining_rejects_heldout_identity_overlap(key):
    from modern_pca.gleason import assert_training_isolated
    a=dict(dataset='sicap',split='train',sample_id='a',file_sha256='aa',source_wsi='slide-a',patient_id='p-a')
    b=dict(dataset='sicap',split='test',sample_id='b',file_sha256='bb',source_wsi='slide-b',patient_id='p-b')
    assert_training_isolated([a],[b])
    b[key]=a[key]
    with pytest.raises(ValueError,match='overlaps held-out'): assert_training_isolated([a],[b])


def test_mined_labels_are_not_ground_truth_or_combined_scores():
    from modern_pca.evaluate_gleason import mining_records
    base=dict(dataset='sicap',split='train',region_role='mixed',region_primary=3,region_secondary=4,
              label=0,class_name='GP3',label_source='published_patch_class')
    predictions,retained=mining_records([base,base],[[.01,.02,.97],[.95,.03,.02]],'teacherhash')
    assert len(predictions)==2 and len(retained)==1
    assert retained[0]['label']==2 and retained[0]['class_name']=='GP5'
    assert retained[0]['annotation_label']==0
    assert retained[0]['mining_model_sha256']=='teacherhash'
    assert base['label']==0  # input annotations must not be modified
    with pytest.raises(ValueError,match='training rows'):
        mining_records([dict(base,split='validation')],[[1,0,0]],'hash')


def mining_fixture(tmp_path):
    root=tmp_path/'manifests'; root.mkdir()
    hashes={}
    for dataset in ('crowd','sicap'):
        for split in ('train','validation','test'):
            rows=[]
            for i in range(3):
                ident=f'{dataset}-{split}-{i}'
                rows.append(dict(dataset=dataset,split=split,sample_id=ident,file_sha256=ident,
                                 source_wsi=ident,patient_id=ident,label=i,class_name=CLASSES[i],
                                 label_source='published_patch_class',region_role='pure',
                                 region_primary=i+3,region_secondary=i+3,cache_file=ident+'.png',cache_sha256=ident))
            if dataset=='sicap' and split=='train':
                rows.append(dict(rows[0],sample_id='candidate',file_sha256='candidate',source_wsi='candidate',
                                 patient_id='candidate',region_role='mixed',region_primary=3,region_secondary=4))
            path=root/f'{dataset}_{split}.csv'; write_csv(path,rows); hashes[path.name]=sha(path)
    prepared=dict(sample_only=False,manifests=hashes,region_grade_sha256='gradehash')
    write_json(tmp_path/'prepared.json',prepared)
    meta=dict(prepared,complete=True,smoke=False,training_cohort='pure',model_sha256='modelhash')
    return root,prepared,meta


def test_mining_candidates_strict_teacher_and_frozen_manifests(tmp_path):
    from modern_pca.evaluate_gleason import mining_candidates
    root,prepared,meta=mining_fixture(tmp_path)
    rows=mining_candidates(root,prepared,meta)
    assert [r['sample_id'] for r in rows]==['candidate']
    for overrides in (dict(smoke=True),dict(training_cohort='published'),dict(complete=False),
                      dict(manifests={}),dict(region_grade_sha256='changed')):
        with pytest.raises(ValueError): mining_candidates(root,prepared,dict(meta,**overrides))
    with pytest.raises(ValueError): mining_candidates(root,dict(prepared,sample_only=True),meta)


@pytest.mark.parametrize('accepted',[True,False])
def test_mining_output_pipeline_with_fake_inference(tmp_path,monkeypatch,accepted):
    from argparse import Namespace
    from modern_pca import evaluate_gleason as module
    root,prepared,meta=mining_fixture(tmp_path)
    model_dir=tmp_path/'model'; model_dir.mkdir(); write_json(model_dir/'run.json',meta)
    class FakeModel:
        def predict(self, rows, verbose):
            assert [r['sample_id'] for r in rows]==['candidate']
            return np.array([[.01,.98,.01] if accepted else [.95,.03,.02]])
    monkeypatch.setattr(module,'gpu_setup',lambda seed:None)
    monkeypatch.setattr(module,'load_model',lambda path:(FakeModel(),meta))
    monkeypatch.setattr(module,'verify_cache',lambda cache,rows:None)
    monkeypatch.setattr(module,'dataset',lambda rows,cache,batch:rows)
    output=tmp_path/'mined'
    module.mine(Namespace(model_dir=model_dir,manifests=root,cache=tmp_path,output=output,batch_size=1),prepared)
    result=json.loads((output/'mining.json').read_text())
    assert result['complete'] and result['candidates']==1 and result['retained']==int(accepted)
    assert result['artifacts']['mined_train.csv']==sha(output/'mined_train.csv')
    with (output/'mined_train.csv').open() as stream:
        reader=csv.DictReader(stream); rows=list(reader)
        assert 'mining_model_sha256' in reader.fieldnames
    assert len(rows)==int(accepted)


def test_release_manifest_generated_once_and_deleted_files_excluded(tmp_path):
    import shutil,subprocess,sys
    root=tmp_path/'repo'; (root/'tools').mkdir(parents=True)
    script=Path(__file__).parents[1]/'tools/package_gleason_release.py'
    shutil.copyfile(script,root/'tools/package_gleason_release.py')
    (root/'RELEASE.json').write_text('{}')
    (root/'deleted.txt').write_text('old')
    subprocess.run(['git','init','-q',str(root)],check=True)
    subprocess.run(['git','-C',str(root),'add','.'],check=True)
    subprocess.run(['git','-C',str(root),'-c','user.name=Test','-c','user.email=test@example.invalid',
                    'commit','-qm','fixture'],check=True)
    (root/'deleted.txt').unlink()
    output=tmp_path/'release.zip'
    subprocess.run([sys.executable,str(root/'tools/package_gleason_release.py'),'--output',str(output)],check=True)
    with zipfile.ZipFile(output) as z:
        assert z.namelist().count('dsa5206-proj/RELEASE.json')==1
        assert 'dsa5206-proj/deleted.txt' not in z.namelist()
        assert 'RELEASE.json' not in json.loads(z.read('dsa5206-proj/RELEASE.json'))['files']
